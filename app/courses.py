"""Courses: a home for everything about one subject — imported slides / readings and a glossary of
the subject's terms. The glossary is what keeps recognition, translation and summaries right about
the course's own vocabulary.

Stored as <sessions dir>/_课程/<course>/course.json, with imported files copied next to it.
course = {name, created, terms: [{term, zh, en, source, added}], materials: [{file, title, pages, added}]}
"""
import difflib
import json
import re
import shutil
import threading
import time
from pathlib import Path

DIR = "_课程"
_lock = threading.Lock()


def _safe(name):
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip() or "未命名课程"


def course_dir(root, name) -> Path:
    return Path(root) / DIR / _safe(name)


def load(root, name):
    p = course_dir(root, name) / "course.json"
    with _lock:
        if p.exists():
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
                d.setdefault("terms", [])
                d.setdefault("materials", [])
                return d
            except Exception:
                pass
    return {"name": name, "created": time.strftime("%Y-%m-%dT%H:%M:%S"), "terms": [], "materials": []}


def save(root, d):
    folder = course_dir(root, d["name"])
    folder.mkdir(parents=True, exist_ok=True)
    with _lock:
        tmp = folder / "course.json.tmp"
        tmp.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(folder / "course.json")
        _export_md(folder, d)
    return d


def mutate(root, name, fn):
    d = load(root, name)
    r = fn(d)
    save(root, d)
    return r if r is not None else d


def names(root):
    base = Path(root) / DIR
    out = []
    if base.exists():
        for f in base.glob("*/course.json"):
            try:
                out.append(json.loads(f.read_text(encoding="utf-8")).get("name") or f.parent.name)
            except Exception:
                pass
    return out


def rename(root, old, new):
    a, b = course_dir(root, old), course_dir(root, new)
    if a.exists() and not b.exists():
        a.rename(b)
    d = load(root, new)
    d["name"] = new
    save(root, d)


def _export_md(folder, d):
    lines = [f"# {d['name']} · 术语表", ""]
    for t in sorted(d["terms"], key=lambda x: x["term"].lower()):
        lines.append(f"- **{t['term']}** — {t.get('zh', '')}" + (f"：{t['en']}" if t.get("en") else ""))
    (folder / "术语表.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------- reading slides / documents ----------------

def extract_text(path) -> list:
    """[(page or slide number, text)] from a PDF, PowerPoint or text file."""
    p = Path(path)
    ext = p.suffix.lower()
    pages = []
    if ext == ".pdf":
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(str(p)).pages):
            try:
                pages.append((i + 1, page.extract_text() or ""))
            except Exception:
                pages.append((i + 1, ""))
    elif ext in (".pptx",):
        from pptx import Presentation
        for i, slide in enumerate(Presentation(str(p)).slides):
            texts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    texts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        texts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide:
                texts.append(slide.notes_slide.notes_text_frame.text)
            pages.append((i + 1, "\n".join(texts)))
    elif ext in (".txt", ".md"):
        pages.append((1, p.read_text(encoding="utf-8", errors="ignore")))
    else:
        raise ValueError("只支持 PDF、PowerPoint（.pptx）和文本文件")
    return [(n, re.sub(r"[ \t]+", " ", t).strip()) for n, t in pages]


# ---------------- terms ----------------

TERM_PROMPT = (
    "你在帮一位用英语上课的中国留学生整理一门课的专业术语表。下面是这门课的课件/资料文字。\n"
    "请找出这门课的**学科专业术语**（概念、定律、模型、方法、专有名词、缩写），不要日常词汇或普通形容词。"
    "按重要程度最多列 {n} 个。每个术语一行，严格用这个格式：\n"
    "术语 | 标准中文译名 | 一句简明英文定义（在这门课里的意思，20 词以内）\n"
    "不要编号，不要其他内容。"
)


def parse_terms(text):
    out = []
    for line in text.splitlines():
        parts = [x.strip().strip("*`") for x in line.strip().lstrip("-•* ").split("|")]
        if len(parts) >= 2 and parts[0] and len(parts[0]) <= 60 and re.search(r"[A-Za-z]", parts[0]):
            out.append({"term": parts[0], "zh": parts[1], "en": parts[2] if len(parts) > 2 else ""})
    return out


def ai_terms(api_key, model, course, text, usage=None, n=50):
    from cloud import claude_chat
    body = text[:40000]
    out = claude_chat(api_key, model, TERM_PROMPT.format(n=n),
                      f"课程：{course}\n\n资料：\n{body}", max_tokens=4000, timeout=120, usage=usage)
    return parse_terms(out)


_STOP = set("""a an the and or of to in on for with by from as at is are was were be been it this that these those
we you they he she i our your their its not no yes can will would should could may might must do does did have has had
which what when where who whom how why than then so if else also more most less least very just only into over under
about after before between each other such some any all both one two three first second new use used using example
slide slides page chapter lecture class today week figure table note notes question questions answer
student students remind raising increases decreases method""".split())


def guess_terms(text, n=30):
    """Without a key: frequent multi-word phrases and capitalised terms (to be checked by the user)."""
    words = re.findall(r"[A-Za-z][A-Za-z\-']+", text)
    low = [w.lower() for w in words]
    counts = {}
    for k in (3, 2):
        for i in range(len(low) - k + 1):
            g = low[i:i + k]
            if g[0] in _STOP or g[-1] in _STOP or any(len(x) < 3 for x in g):
                continue
            key = " ".join(g)
            counts[key] = counts.get(key, 0) + 1
    for w in words:
        if (w.isupper() and 2 <= len(w) <= 6) or (w[0].isupper() and w.lower() not in _STOP and len(w) > 4):
            counts[w] = counts.get(w, 0) + 0.5
    best = sorted(((c, t) for t, c in counts.items() if c >= 2), reverse=True)
    out, seen = [], set()
    for c, t in best:
        if any(t in s or s in t for s in seen):
            continue
        seen.add(t)
        out.append({"term": t, "zh": "", "en": ""})
        if len(out) >= n:
            break
    return out


def merge_terms(existing, new, source):
    have = {t["term"].lower(): t for t in existing}
    added = 0
    for t in new:
        k = t["term"].lower()
        if k in have:
            old = have[k]
            if not old.get("zh") and t.get("zh"):
                old["zh"] = t["zh"]
            if not old.get("en") and t.get("en"):
                old["en"] = t["en"]
            continue
        t = dict(t, source=source, added=time.strftime("%Y-%m-%d"))
        existing.append(t)
        have[k] = t
        added += 1
    return added


# ---------------- using the glossary ----------------

def _wordset(s):
    return re.findall(r"[a-z][a-z\-']+", s.lower())


def relevant(terms, text, limit=12):
    """Terms that appear in `text` — or almost appear (a misheard 'elastisity'), so proofreading
    can put the right spelling back. Keeps prompts short: only these are sent with a sentence."""
    if not terms or not text:
        return []
    words = _wordset(text)
    low = " " + " ".join(words) + " "
    cands = words + [a + b for a, b in zip(words, words[1:])]   # "dead weight" ~ "deadweight"
    hits = []
    for t in terms:
        tw = _wordset(t["term"])
        if not tw:
            continue
        if " " + " ".join(tw) + " " in low:
            hits.append((2.0, t))
            continue
        # fuzzy: every long word of the term has a close match in the sentence
        long = [w for w in tw if len(w) >= 5]
        if long and all(any(difflib.SequenceMatcher(None, w, x).ratio() >= 0.8 for x in cands if abs(len(x) - len(w)) <= 3)
                        for w in long):
            hits.append((1.0, t))
    hits.sort(key=lambda x: -x[0])
    return [t for _, t in hits[:limit]]


def glossary_text(terms, limit=200):
    lines = []
    for t in terms[:limit]:
        s = t["term"]
        if t.get("zh"):
            s += " — " + t["zh"]
        lines.append(s)
    return "\n".join(lines)


def copy_material(root, name, path):
    folder = course_dir(root, name) / "课件"
    folder.mkdir(parents=True, exist_ok=True)
    dst = folder / Path(path).name
    if Path(path).resolve() != dst.resolve():
        shutil.copy2(path, dst)
    return dst
